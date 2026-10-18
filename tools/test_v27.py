# -*- coding: utf-8 -*-
"""
v2.7 新增能力的单元测试（不启动界面）：
    1. 跨批次台账（dup_history）：记 / 查 / 金额不同不算 / 同文件夹不算 / 清空
    2. 报销清单 write_pack_list：行数、合计、号码是文本格式
    3. 红字票 ↔ 蓝票匹配（_compute_red_match）
    4. 「撤销上次」只撤最近一批（v2.7 语义修正）
    5. 流式扫描的 on_item 回调（parse_folder）
    6. 视图偏好 save_view 落盘

运行: python tools/test_v27.py
"""
import glob
import json
import os
import shutil
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import dup_history as DH          # noqa: E402
import export_core as EC          # noqa: E402
import invoice_parser as IP       # noqa: E402
import rename_core as RC          # noqa: E402
import bridge                     # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TESTS = os.path.join(ROOT, "tests")

fails = []


def check(cond, msg, extra=None):
    print(("  ✓ " if cond else "  ✗ ") + msg
          + ("" if extra is None else "  → " + str(extra)))
    if not cond:
        fails.append(msg)


def main():
    tmp = tempfile.mkdtemp(prefix="v27_")

    # ---------- 1. 跨批次台账 ----------
    print("─" * 70)
    print("1) 跨批次台账 dup_history")
    idx = DH.HistoryIndex(os.path.join(tmp, "hist.json"))
    n = idx.record([
        {"no": "253120000001", "amount": 100.0, "date": "2026-09-12",
         "folder": os.path.join(tmp, "9月"), "file": "a.pdf"},
        {"no": "253120000002", "amount": 200.0, "date": "2026-09-13",
         "folder": os.path.join(tmp, "9月"), "file": "b.pdf"},
    ])
    check(n == 2, "记下 2 张票", n)
    hit = idx.lookup("253120000001", 100.0, os.path.join(tmp, "10月"))
    check(hit is not None, "换到别的文件夹能查出来（这就是跨批次）")
    check(
        idx.lookup("253120000001", 999.0, os.path.join(tmp, "10月")) is None,
        "金额不一致不报（红冲重开是正常业务）")
    check(
        idx.lookup("253120000001", 100.0, os.path.join(tmp, "9月")) is None,
        "同一个文件夹再打开一次不报")
    check(idx.lookup("不存在的票号", 1.0, tmp) is None, "没记过的票号查不到")
    check(idx.count() == 2, "台账里 2 条", idx.count())
    idx.clear()
    check(idx.count() == 0 and idx.lookup("253120000001", 100.0, tmp) is None,
          "清空后查不到（操作历史里的「清空记忆」）")

    # ---------- 2. 报销清单 ----------
    print("\n" + "─" * 70)
    print("2) 报销清单 write_pack_list")
    xp = os.path.join(tmp, "报销清单.xlsx")
    rows = [
        {"name": "上海某某—94222.21.pdf", "buyer": "上海某某实业有限公司",
         "amount": 94222.21, "no": "253120000002764865"},
        {"name": "哈啰出行_行程单.pdf", "buyer": "（附件）",
         "amount": 1410.30, "no": ""},
        {"name": "肉界食品—77901.51.pdf", "buyer": "肉界食品（苏州）有限公司",
         "amount": 77901.51, "no": "253120000002779451"},
    ]
    r = EC.write_pack_list(xp, rows, {"title": "报销单_20261018"})
    check(r["ok"] and os.path.exists(xp), "生成了 xlsx")
    check(abs(r["total"] - 173534.02) < 0.01, "合计算对", r["total"])
    from openpyxl import load_workbook
    wb = load_workbook(xp)
    ws = wb.active
    check(ws["A1"].value == "报销单_20261018", "抬头写的是报销单名", ws["A1"].value)
    check(ws["A2"].value == "序号", "第二行是表头", ws["A2"].value)
    check(ws["E3"].value == "253120000002764865", "票号原样写入", ws["E3"].value)
    check(ws["E3"].number_format == "@", "票号列是文本格式（不会变科学计数法）",
          ws["E3"].number_format)
    # 合计行在第 3+3 = 6 行
    check(ws["B6"].value == "合计 3 项", "末尾有合计行", ws["B6"].value)
    check(abs((ws["D6"].value or 0) - 173534.02) < 0.01, "合计单元格数值正确",
          ws["D6"].value)

    # ---------- 3. 红字票 ↔ 蓝票 ----------
    print("\n" + "─" * 70)
    print("3) 红字票关联蓝票 _compute_red_match")
    api = bridge.Api()
    api.items = []

    def mk(name, amount, seller, date, red=False):
        it = IP.InvoiceInfo(os.path.join(tmp, name))
        it.file_name = name
        it.kind = "invoice"
        it.ok = True
        it.amount = amount
        it.seller = seller
        it.date = date
        it.is_red = red
        it.invoice_no = "NO" + name
        return it

    api.items = [
        mk("蓝票.pdf", 2180.0, "广州某某物流有限公司", "2026-08-12"),
        mk("红票.pdf", -2180.0, "广州某某物流有限公司", "2026-08-18", red=True),
        mk("无关蓝票.pdf", 500.0, "别的公司", "2026-08-12"),
        mk("别家红票.pdf", -500.0, "某贸易公司", "2026-08-12", red=True),
    ]
    api._compute_red_match()
    m = api.red_match.get(1)
    check(m is not None and m["i"] == 0, "红票匹配到了同销售方、金额相反的蓝票",
          m and m["file"])
    check(api.red_by.get(0, {}).get("i") == 1, "蓝票也记下了是谁冲的")
    check(3 not in api.red_match, "销售方不同不硬配（别家红票没配上）")

    # ---------- 4. 撤销只撤最近一批 ----------
    print("\n" + "─" * 70)
    print("4) 「撤销上次」只撤最近一批")
    work = os.path.join(tmp, "撤销测试")
    os.makedirs(work, exist_ok=True)
    srcs = []
    for i in range(3):
        p = os.path.join(work, "orig%d.pdf" % i)
        with open(p, "wb") as f:
            f.write(b"%PDF-1.4\n% test\n")
        srcs.append(p)

    def run_batch(pairs, name):
        entries = [RC.Entry(src, base, "", "", None, True) for src, base in pairs]
        plan = RC.build_plan(entries, copy_mode=False, dedup="index",
                             distribute=False)
        rn = RC.Runner(plan, do_backup=False, backup_root=work)
        rn.run()
        RC.append_log(work, rn.log)
        return rn

    run_batch([(srcs[0], "第一批0"), (srcs[1], "第一批1")], "第一批")
    # 第二批只改一个
    run_batch([(os.path.join(work, "第一批0.pdf"), "第二批0")], "第二批")

    api2 = bridge.Api()
    api2.folder = work
    api2.items = []
    check(api2.undo_count() == 1, "「撤销上次」只算最近一批（1 条）", api2.undo_count())
    api2.undo()
    for _ in range(100):
        if not api2.task["busy"]:
            break
        import time
        time.sleep(0.05)
    now = sorted(os.listdir(work))
    check("第一批0.pdf" in now, "最近一批被还原回上一版名字", now)
    check("第一批1.pdf" in now, "**更早那批没动**（老版本会把两批一起撤）", now)

    # ---------- 4b. 打包带清单（真实复制 + xlsx） ----------
    print("\n" + "─" * 70)
    print("4b) 打包文件并生成报销清单")
    api4 = bridge.Api()
    api4.folder = TESTS
    api4.items = IP.parse_folder(TESTS)[:3]
    api4.checked = [True] * len(api4.items)
    outdir = os.path.join(tmp, "打包输出")
    os.makedirs(outdir, exist_ok=True)
    api4.pack_reimbursement([0, 1], outdir, "测试报销", True)
    for _ in range(200):
        if not api4.task["busy"]:
            break
        import time
        time.sleep(0.05)
    pay = api4.task.get("payload") or {}
    check(pay.get("ok") is True, "打包任务成功", pay.get("msg") or pay.get("dir"))
    got = sorted(os.listdir(pay.get("dir") or outdir)) if pay.get("dir") else []
    check(any(f.endswith(".xlsx") for f in got), "包里生成了清单 xlsx", got)
    check(sum(1 for f in got if not f.endswith(".xlsx")) == 2, "两张票都复制进去了", got)
    check(pay.get("list") == "报销清单.xlsx", "返回里写明了清单文件名", pay.get("list"))

    # ---------- 5. 流式扫描回调 ----------
    print("\n" + "─" * 70)
    print("5) 流式扫描 on_item")
    seen = []
    out = IP.parse_folder(TESTS, on_item=lambda i, it, total: seen.append((i, it.file_name, total)))
    check(len(seen) == len(out), "每张票都回调了一次", "%d / %d" % (len(seen), len(out)))
    check(len(seen) > 0 and seen[0][2] == len(out), "回调里带的 total 正确",
          seen[0][2] if seen else None)

    # ---------- 5b. 流式扫描：扫描途中就推中间结果 ----------
    print("\n" + "─" * 70)
    print("5b) 流式扫描（bridge._live_on_item / poll.partial）")
    big = os.path.join(tmp, "大批量")
    os.makedirs(big, exist_ok=True)
    import shutil as _sh
    src_one = sorted(glob.glob(os.path.join(TESTS, "*.pdf")))[0]
    for i in range(48):
        _sh.copyfile(src_one, os.path.join(big, "f%03d.pdf" % i))
    api5 = bridge.Api()
    api5.folder = big
    api5.rescan()
    parts, counts = 0, []
    for _ in range(400):
        t = api5.poll()
        if t.get("partial"):
            parts += 1
            counts.append(len(t["partial"]["rows"]))
        if not t.get("busy"):
            break
        import time
        time.sleep(0.01)
    check(parts >= 1, "扫描途中推送了中间结果", "%d 次" % parts)
    check(counts and counts[0] < 48, "第一批小于全量（确实是边读边出）",
          (counts[:4] if counts else None))
    check(len(api5.items) == 48, "最终仍是完整 48 张", len(api5.items))

    # ---------- 6. 视图偏好 ----------
    print("\n" + "─" * 70)
    print("6) 视图偏好 save_view")
    api3 = bridge.Api()
    old = dict(api3.cfg.get("view") or {})
    r = api3.save_view({"sort": "amount", "folded": False})
    check(r["ok"] and r["view"]["sort"] == "amount", "save_view 返回新视图", r["view"])
    check(api3.cfg["view"]["folded"] is False, "折叠状态记进了配置")
    # 还原，别把测试偏好留给真机（配置文件是真实的那一份）
    api3.save_view(old or {"sort": "file", "folded": True})

    shutil.rmtree(tmp, ignore_errors=True)

    print("\n" + "=" * 70)
    if fails:
        print("未通过 %d 项：" % len(fails))
        for f in fails:
            print("  - " + f)
        sys.exit(1)
    print("v2.7 新增能力测试全部通过 ✅")


if __name__ == "__main__":
    main()
